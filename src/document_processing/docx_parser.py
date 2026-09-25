"""
Pipeline 2 - Document Extraction & Parsing Module (DOCX)
Deterministic text and section parser for DOCX documents.
Extracts headings, paragraph references, tables, and document metadata.
"""

import os
import re
import logging
from dataclasses import dataclass, field
from typing import List, Dict, Any, Optional

logger = logging.getLogger("DOCXParser")


@dataclass
class ParsedSection:
    section_id: str
    heading: str
    content: str
    paragraph_ref: str
    location_ref: str


@dataclass
class ParsedDocxDocument:
    doc_id: str
    title: str
    doc_type: str  # POLICY, SOP, FAQ, etc.
    version: str
    effective_date: str
    file_path: str
    sections: List[ParsedSection] = field(default_factory=list)
    metadata: Dict[str, Any] = field(default_factory=dict)
    full_text: str = ""


class DOCXParser:
    """
    Deterministic DOCX Parser using python-docx with fallback text parsing.
    Extracts sections, headings, paragraph tags, doc_id, and version metadata.
    """

    DOC_ID_PATTERN = re.compile(r'\b(POL|SOP|FAQ|MAN|MANUAL|GUIDE|HANDBOOK)-[A-Z0-9]+(?:-[0-9]+)?\b', re.IGNORECASE)
    SECTION_HEADER_PATTERN = re.compile(r'^(?:Section\s+[0-9.]+|[0-9]+\.[0-9.]*\s+[A-Z].*|Heading\s+[0-9]+)', re.IGNORECASE)

    def parse(self, file_path: str, doc_id_override: Optional[str] = None) -> ParsedDocxDocument:
        # FIX: Guard against null or missing file path
        if not file_path or not os.path.exists(file_path):
            raise FileNotFoundError(f"DOCX file not found or invalid: {file_path}")

        file_name = os.path.basename(file_path)
        sections: List[ParsedSection] = []
        raw_paragraphs: List[str] = []

        try:
            import docx  # type: ignore
            doc = docx.Document(file_path)
            current_heading = "General"
            current_section_id = "SEC-01"
            sec_idx = 1
            para_counter = 0
            current_text_buf: List[str] = []

            for p in doc.paragraphs:
                p_text = p.text.strip()
                if not p_text:
                    continue
                para_counter += 1
                raw_paragraphs.append(p_text)

                is_heading = (
                    (hasattr(p, "style") and p.style and p.style.name and p.style.name.startswith("Heading")) or
                    bool(self.SECTION_HEADER_PATTERN.match(p_text))
                )
                if is_heading:
                    if current_text_buf:
                        sections.append(ParsedSection(
                            section_id=current_section_id,
                            heading=current_heading,
                            content="\n".join(current_text_buf),
                            paragraph_ref=f"Para {para_counter - len(current_text_buf)}-{para_counter - 1}",
                            location_ref=f"{current_section_id} ({current_heading})"
                        ))
                        current_text_buf = []
                    current_heading = p_text
                    sec_idx += 1
                    current_section_id = f"SEC-{sec_idx:02d}"
                else:
                    current_text_buf.append(p_text)

            # Flush remaining buffer
            if current_text_buf:
                sections.append(ParsedSection(
                    section_id=current_section_id,
                    heading=current_heading,
                    content="\n".join(current_text_buf),
                    paragraph_ref=f"Para {para_counter - len(current_text_buf) + 1}-{para_counter}",
                    location_ref=f"{current_section_id} ({current_heading})"
                ))

            # Extract tables if present
            for t_idx, table in enumerate(doc.tables, start=1):
                table_lines = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        table_lines.append(" | ".join(row_data))
                if table_lines:
                    t_text = "\n".join(table_lines)
                    raw_paragraphs.append(t_text)
                    sections.append(ParsedSection(
                        section_id=f"TABLE-{t_idx}",
                        heading=f"Table {t_idx}",
                        content=t_text,
                        paragraph_ref=f"Table {t_idx}",
                        location_ref=f"Table {t_idx}"
                    ))

        except Exception as e:
            logger.debug(f"python-docx parsing failed on {file_path}, falling back to plain text read: {e}")
            # Fallback for plain text files or mock docx files
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read().strip()
                raw_paragraphs = [content]
                sections.append(ParsedSection(
                    section_id="SEC-01",
                    heading="Document Body",
                    content=content,
                    paragraph_ref="Para 1-End",
                    location_ref="SEC-01 (Document Body)"
                ))
            except Exception as read_err:
                logger.error(f"Failed fallback reading {file_path}: {read_err}")
                raw_paragraphs = [""]
                sections.append(ParsedSection(
                    section_id="SEC-01",
                    heading="Empty Document",
                    content="",
                    paragraph_ref="Para 0",
                    location_ref="SEC-01"
                ))

        full_text = "\n\n".join(raw_paragraphs)
        metadata = self._extract_metadata(file_name, full_text, doc_id_override)

        return ParsedDocxDocument(
            doc_id=metadata["doc_id"],
            title=metadata["title"],
            doc_type=metadata["doc_type"],
            version=metadata["version"],
            effective_date=metadata["effective_date"],
            file_path=file_path,
            sections=sections,
            metadata=metadata,
            full_text=full_text
        )

    def _extract_metadata(self, file_name: str, content: str, doc_id_override: Optional[str] = None) -> Dict[str, Any]:
        doc_id = doc_id_override or ""
        if not doc_id:
            match = self.DOC_ID_PATTERN.search(file_name) or self.DOC_ID_PATTERN.search(content[:1000])
            if match:
                doc_id = match.group(0).upper()
            else:
                doc_id = os.path.splitext(file_name)[0].upper()

        doc_type = "POLICY"
        if doc_id.startswith("SOP"):
            doc_type = "SOP"
        elif doc_id.startswith("FAQ"):
            doc_type = "FAQ"

        # FIX: Robust regex capturing versions preceded by underscores, hyphens, or spaces
        version = "v1.0"
        v_match = (
            re.search(r'(?:[_\s-]|^)v?(\d+(?:\.\d+)+)', file_name, re.IGNORECASE) or
            re.search(r'Version\s*:?\s*v?(\d+(?:\.\d+)*)', content[:1000], re.IGNORECASE) or
            re.search(r'(?:[_\s-]|^)v(\d+(?:\.\d+)*)', content[:1000], re.IGNORECASE)
        )
        if v_match:
            v_val = v_match.group(1)
            if v_val:
                version = f"v{v_val}" if not v_val.startswith("v") else v_val

        effective_date = "2026-01-01"
        date_match = re.search(
            r'(?:Effective Date|Date)\s*:?\s*([0-9]{4}-[0-9]{2}-[0-9]{2})',
            content[:1000],
            re.IGNORECASE
        )
        if date_match:
            effective_date = date_match.group(1)

        title = file_name
        lines = [line.strip() for line in content.splitlines() if line.strip()]
        if lines:
            first_line = lines[0]
            if 5 < len(first_line) < 120 and not first_line.startswith("PK"):
                title = first_line

        return {
            "doc_id": doc_id,
            "title": title,
            "doc_type": doc_type,
            "version": version,
            "effective_date": effective_date,
            "file_name": file_name
        }
